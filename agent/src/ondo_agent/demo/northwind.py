"""The Northwind client drive from the design, as real files.

Sample data only, kept consistent with the mockups: twelve signed contracts, last
quarter's pack, a Q3 workbook whose Halleck Logistics row (row 14) carries last
year's 3.5% uplift while the signed contract says 5% from 1 July, and a payroll
workbook the policy excludes. The six changed accounts add 41,190 to total annual
value, which is the figure on the approval card.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class Account:
    name: str
    file: str
    current: int
    uplift_pct: float  # what the signed contract says
    renewal: str
    workbook_pct: float | None = None  # what the workbook says, when it disagrees

    @property
    def new_value(self) -> int:
        return round(self.current * (1 + self.uplift_pct / 100))


ACCOUNTS = [
    Account("Ashgrove Cold Chain", "Ashgrove_MSA_2025.docx", 210_000, 4.0, "2026-07-01"),
    Account("Brightwater Ports", "Brightwater_Ports_Agreement.pdf", 156_000, 4.5, "2026-08-01"),
    Account("Corran Haulage", "Corran_Haulage_MSA.docx", 128_500, 4.0, "2026-07-15"),
    Account("Dunmore Warehousing", "Dunmore_Warehousing_Contract.pdf", 172_900, 5.0, "2026-09-01"),
    Account("Eastfield Rail", "Eastfield_Rail_MSA.docx", 98_000, 0.0, "2027-01-01"),
    Account("Fenwick Maritime", "Fenwick_Maritime_Agreement.pdf", 143_250, 0.0, "2027-03-01"),
    Account("Garston Supply", "Garston_Supply_SOW.docx", 61_400, 0.0, "2027-02-01"),
    Account("Harrow Parcel", "Harrow_Parcel_MSA.docx", 77_800, 0.0, "2026-12-01"),
    Account("Ivel Distribution", "Ivel_Distribution_Contract.pdf", 119_600, 0.0, "2027-04-01"),
    Account("Jarrow Trucking", "Jarrow_Trucking_MSA.docx", 88_100, 0.0, "2027-01-15"),
    Account("Pemberton Freight", "Pemberton_Freight_SOW.docx", 92_000, 3.0, "2026-07-01"),
    Account("Halleck Logistics", "Halleck_MSA_2026.pdf", 184_500, 5.0, "2026-07-01", workbook_pct=3.5),
]

TOTAL_INCREASE = sum(a.new_value - a.current for a in ACCOUNTS)  # 41,190


def _contract_lines(a: Account) -> list[str]:
    uplift = (
        f"7.2 Annual uplift. From {a.renewal} the annual fee increases by {a.uplift_pct:g}% to {a.new_value:,} GBP."
        if a.uplift_pct
        else f"7.2 Annual uplift. No uplift applies at the next renewal on {a.renewal}. The annual fee remains {a.current:,} GBP."
    )
    return [
        f"MASTER SERVICES AGREEMENT - {a.name.upper()}",
        "Between Northwind Operations Ltd and " + a.name + ".",
        "1. Services. Freight, storage and handling services as set out in Schedule 1.",
        f"6. Fees. The annual fee for the current term is {a.current:,} GBP, invoiced quarterly.",
        f"7.1 Renewal. This agreement renews automatically on {a.renewal} for a further twelve months.",
        uplift,
        "7.3 Notice. Either party may give ninety days' written notice before renewal.",
        "Signed for and on behalf of both parties.",
    ]


def write_pdf(path: Path, lines: list[str]) -> None:
    """A minimal text PDF, enough for pypdf to extract. No dependency needed."""

    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    pages = [lines[i : i + 48] for i in range(0, len(lines), 48)] or [[]]
    objs: list[bytes] = []
    n_pages = len(pages)
    # 1 catalog, 2 pages, 3 font, then (page, content) pairs
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n_pages))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, pl in enumerate(pages):
        body = "BT /F1 10 Tf 60 780 Td 15 TL " + " ".join(f"({esc(line)}) Tj T*" for line in pl) + " ET"
        stream = body.encode("latin-1", errors="replace")
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {5 + 2 * i} 0 R >>".encode()
        )
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def _write_docx(path: Path, lines: list[str]) -> None:
    import docx

    d = docx.Document()
    d.add_heading(lines[0], level=1)
    for line in lines[1:]:
        d.add_paragraph(line)
    d.save(str(path))


def build(root: Path) -> Path:
    """Write the drive under ``root/Northwind client drive`` and return that folder."""
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    drive = Path(root) / "Northwind client drive"
    contracts = drive / "Contracts"
    contracts.mkdir(parents=True, exist_ok=True)
    (drive / "HR").mkdir(exist_ok=True)

    for a in ACCOUNTS:
        p = contracts / a.file
        if p.suffix == ".pdf":
            write_pdf(p, _contract_lines(a))
        else:
            _write_docx(p, _contract_lines(a))

    wb = Workbook()
    ws = wb.active
    ws.title = "Renewals"
    ws.append(["Northwind Q3 renewals", None, None, None, None, None])
    ws.append(["Account", "Contract", "Renewal date", "Current annual value", "Uplift %", "New annual value"])
    for i, a in enumerate(ACCOUNTS, start=3):
        pct = a.workbook_pct if a.workbook_pct is not None else a.uplift_pct
        ws.append([a.name, a.file, a.renewal, a.current, pct, f"=ROUND(D{i}*(1+E{i}/100),0)"])
    last = 2 + len(ACCOUNTS)
    ws.append(["Total", None, None, f"=SUM(D3:D{last})", None, f"=SUM(F3:F{last})"])
    wb.save(drive / "Q3_Renewals.xlsx")

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Q2 renewal pack — Northwind"
    rows = [["Account", "Annual value"]] + [[a.name, f"{a.current:,}"] for a in ACCOUNTS[:6]]
    tbl = s.shapes.add_table(len(rows), 2, Inches(0.5), Inches(1.5), Inches(8), Inches(4)).table
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            tbl.cell(r, c).text = v
    prs.save(drive / "Q2_Renewal_Pack_final.pptx")

    pay = Workbook()
    pay.active.append(["Employee", "Salary"])
    pay.active.append(["Redacted", 0])
    pay.save(drive / "HR" / "Payroll_2026_confidential.xlsx")

    (drive / "Billing_export_2026-09.csv").write_text(
        "account,annual_value\n" + "\n".join(f"{a.name},{a.current}" for a in ACCOUNTS) + "\n"
    )
    return drive
