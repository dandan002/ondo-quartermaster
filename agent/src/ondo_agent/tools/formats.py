"""Read and write office files as files, never through the application.

``openpyxl``, ``python-docx``, ``python-pptx`` and ``pypdf``: deterministic, no
app install, no visible window, no model involved in the mechanics. The reliable
product is "Ondo edits the workbook and shows you the diff", not "Ondo types in
Excel".
"""

from __future__ import annotations

import csv
import difflib
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

KINDS = {
    ".xlsx": "Workbook",
    ".xlsm": "Workbook",
    ".docx": "Document",
    ".pptx": "Deck",
    ".pdf": "PDF",
    ".csv": "Data",
    ".tsv": "Data",
    ".txt": "Text",
    ".md": "Text",
    ".json": "Data",
}


def kind_of(path: Path) -> str:
    k = KINDS.get(path.suffix.lower())
    if k in ("PDF", "Document") and any(w in path.stem.lower() for w in ("msa", "contract", "agreement", "sow")):
        return "Contract"
    return k or "File"


@dataclass
class ReadOut:
    text: str
    meta: dict[str, Any] = field(default_factory=dict)


def _fmt(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:.10g}"
    return str(v)


def read_xlsx(path: Path, *, sheet: str | None = None, max_rows: int = 300) -> ReadOut:
    from openpyxl import load_workbook

    values = load_workbook(path, data_only=True, read_only=True)
    formulas = load_workbook(path, data_only=False, read_only=True)
    names = [sheet] if sheet else values.sheetnames
    out, meta = [], {"sheets": values.sheetnames}
    for name in names:
        if name not in values.sheetnames:
            out.append(f'Sheet "{name}" not found. Sheets: {", ".join(values.sheetnames)}')
            continue
        ws, wf = values[name], formulas[name]
        out.append(f'## Sheet "{name}" ({ws.max_row} rows x {ws.max_column} columns)')
        for i, (row_v, row_f) in enumerate(zip(ws.iter_rows(), wf.iter_rows(), strict=False)):
            if i >= max_rows:
                out.append(
                    f"... {ws.max_row - max_rows} more rows. Read again with a sheet name and a higher max_rows if needed."
                )
                break
            cells = []
            for cv, cf in zip(row_v, row_f, strict=False):
                v = cv.value
                f = cf.value
                if v is None and isinstance(f, str) and f.startswith("="):
                    cells.append(f"{cv.coordinate}: {f}")
                elif v is not None:
                    s = _fmt(v)
                    if isinstance(f, str) and f.startswith("="):
                        s += f" ({f})"
                    cells.append(f"{cv.coordinate}: {s}")
            if cells:
                out.append(" | ".join(cells))
    values.close()
    formulas.close()
    return ReadOut("\n".join(out), meta)


def read_docx(path: Path) -> ReadOut:
    import docx

    d = docx.Document(str(path))
    out = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            out.append("#" * min(int(level) + 1, 6) + " " + t)
        else:
            out.append(t)
    for ti, table in enumerate(d.tables, 1):
        out.append(f"## Table {ti}")
        for row in table.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return ReadOut("\n".join(out), {"paragraphs": len(d.paragraphs), "tables": len(d.tables)})


def read_pptx(path: Path) -> ReadOut:
    from pptx import Presentation

    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"## Slide {i}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    out.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
    return ReadOut("\n".join(out), {"slides": len(prs.slides)})


def read_pdf(path: Path, *, pages: str | None = None) -> ReadOut:
    from pypdf import PdfReader

    r = PdfReader(str(path))
    n = len(r.pages)
    wanted = _page_range(pages, n)
    out = []
    for i in wanted:
        text = (r.pages[i].extract_text() or "").strip()
        out.append(f"## Page {i + 1}\n{text if text else '[no extractable text: this page may be a scan]'}")
    return ReadOut("\n".join(out), {"pages": n})


def _page_range(spec: str | None, n: int) -> list[int]:
    if not spec:
        return list(range(n))
    out: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            a, b = part.split("-", 1)
            out += list(range(int(a) - 1, min(int(b), n)))
        elif part:
            out.append(int(part) - 1)
    return [i for i in out if 0 <= i < n]


def read_text(path: Path, *, max_rows: int = 300) -> ReadOut:
    raw = path.read_bytes().decode("utf-8", errors="replace")
    lines = raw.splitlines()
    if len(lines) > max_rows and path.suffix.lower() in (".csv", ".tsv"):
        return ReadOut("\n".join(lines[:max_rows]) + f"\n... {len(lines) - max_rows} more rows.", {"lines": len(lines)})
    return ReadOut(raw, {"lines": len(lines)})


def read_any(
    path: Path, *, sheet: str | None = None, pages: str | None = None, max_rows: int = 300, max_chars: int = 60_000
) -> ReadOut:
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        r = read_xlsx(path, sheet=sheet, max_rows=max_rows)
    elif ext == ".docx":
        r = read_docx(path)
    elif ext == ".pptx":
        r = read_pptx(path)
    elif ext == ".pdf":
        r = read_pdf(path, pages=pages)
    else:
        r = read_text(path, max_rows=max_rows)
    if len(r.text) > max_chars:
        r.text = (
            r.text[:max_chars]
            + f"\n... truncated at {max_chars:,} characters. Read a sheet or page range for the rest."
        )
        r.meta["truncated"] = True
    return r


# -- writing ---------------------------------------------------------------------


@dataclass
class CellEdit:
    sheet: str
    cell: str
    value: Any


def _coerce(v: Any) -> Any:
    if isinstance(v, str):
        s = v.strip()
        if s.startswith("="):
            return s
        try:
            return int(s.replace(",", "")) if s.replace(",", "").lstrip("-").isdigit() else float(s.replace(",", ""))
        except ValueError:
            return v
    return v


def plan_workbook_edits(path: Path, edits: list[CellEdit]) -> list[dict[str, Any]]:
    """The diff, computed before anything is written."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=False)
    changes = []
    for e in edits:
        if e.sheet not in wb.sheetnames:
            raise ValueError(f'sheet "{e.sheet}" not in {path.name} (sheets: {", ".join(wb.sheetnames)})')
        before = wb[e.sheet][e.cell].value
        after = _coerce(e.value)
        changes.append(
            {
                "sheet": e.sheet,
                "cell": e.cell.upper(),
                "before": _fmt(before),
                "after": _fmt(after),
                "changed": _fmt(before) != _fmt(after),
            }
        )
    return changes


def apply_workbook_edits(path: Path, edits: list[CellEdit]) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    for e in edits:
        wb[e.sheet][e.cell] = _coerce(e.value)
    wb.save(path)


def create_workbook(path: Path, sheets: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for s in sheets:
        ws = wb.create_sheet(str(s.get("name", "Sheet1"))[:31])
        for row in s.get("rows", []):
            ws.append([_coerce(v) for v in row])
    wb.save(path)


def create_document(path: Path, title: str, paragraphs: list[str], table: list[list[str]] | None = None) -> None:
    import docx

    d = docx.Document()
    if title:
        d.add_heading(title, level=1)
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=0, cols=max(len(r) for r in table))
        t.style = "Table Grid"
        for r in table:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = str(v)
    d.save(str(path))


def unified_diff(before: str, after: str, name: str) -> str:
    return "".join(difflib.unified_diff(before.splitlines(True), after.splitlines(True), f"a/{name}", f"b/{name}"))


def rows_preview(rows: list[list[Any]], limit: int = 15) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    for r in rows[:limit]:
        w.writerow(r)
    more = f"... {len(rows) - limit} more rows\n" if len(rows) > limit else ""
    return buf.getvalue() + more
