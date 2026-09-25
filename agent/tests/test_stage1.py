"""Stage 1: files, properly.

Done when: "build the renewal pack from these twelve contracts" works end to end
with no GUI automation anywhere, and every write shows a diff first.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from ondo_agent import log as L
from ondo_agent.approvals import AutoApprovals
from ondo_agent.demo.northwind import ACCOUNTS, TOTAL_INCREASE
from ondo_agent.demo.policies import renewal_pack_policy
from ondo_agent.models.adapters.scripted import call, say
from ondo_agent.models.gateway import scripted_client
from ondo_agent.permissions import PermissionBroker, Grant, Policy
from ondo_agent.runtime import assemble

from conftest import base_config

REQUEST = ("Build the Q3 renewal pack for Northwind from the contracts in the client folder, "
           "and flag anything that uplifts above five per cent.")


async def test_renewal_pack_end_to_end(drive, tmp_path):
    approvals = AutoApprovals(True, by="mara.okonjo")
    cfg = base_config(drive, tmp_path)
    a = await assemble(cfg, model=scripted_client(renewal_pack_policy(drive)), approvals=approvals)
    res = await a.harness.run(REQUEST)
    assert res.status == "finished", res
    assert f"adding {TOTAL_INCREASE:,}" in res.answer
    assert "Halleck Logistics: the workbook had 3.5%; the signed contract (clause 7.2) says 5%." in res.answer

    events = a.log.events
    # Twelve contracts plus the workbook, read in one parallel turn.
    reads = [e for e in events if e.type == L.FILE_ACCESS and e.data["op"] == "read"]
    assert len(reads) == 13
    # No GUI automation anywhere: only file tools were called.
    assert {e.data["name"] for e in events if e.type == L.TOOL_CALL} <= {"list_folder", "read_file", "edit_workbook", "create_workbook"}

    # Every write: diff first, then the approval, then the edit.
    for tool in ("edit_workbook", "create_workbook"):
        diff = next(e for e in events if e.type == L.DIFF_PROPOSED and e.source == f"tool:{tool}")
        req = next(e for e in events if e.type == L.APPROVAL_REQUESTED and e.data["tool"] == tool)
        edited = next(e for e in events if e.type == L.FILE_ACCESS and e.data["op"] == "edited" and e.data["path"] == diff.data["path"])
        assert diff.seq < req.seq < edited.seq
        assert req.data["diff"] == diff.data["diff"]
    wb_diff = next(e for e in events if e.type == L.DIFF_PROPOSED and e.source == "tool:edit_workbook")
    assert wb_diff.data["diff"] == "Renewals!E14: 3.5 -> 5"
    assert len(approvals.seen) == 2 and all("file_write" in r.effects for r in approvals.seen)

    # The files really changed.
    assert load_workbook(drive / "Q3_Renewals.xlsx")["Renewals"]["E14"].value == 5
    pack = load_workbook(drive / "Q3_Renewal_Pack.xlsx")["Q3 renewal pack"]
    rows = list(pack.iter_rows(values_only=True))
    assert len(rows) == 1 + len(ACCOUNTS)
    assert sum(r[5] for r in rows[1:]) == TOTAL_INCREASE

    # The excluded HR folder was never opened, and the exclusion was logged when it bit.
    assert not any("Payroll" in e.data.get("path", "") for e in reads)
    assert any(e.data["op"] == "excluded" and e.data["path"].endswith("/HR") for e in events if e.type == L.FILE_ACCESS)


async def test_refused_write_changes_nothing(drive, tmp_path):
    before = (drive / "Q3_Renewals.xlsx").read_bytes()
    cfg = base_config(drive, tmp_path)
    a = await assemble(cfg, model=scripted_client(renewal_pack_policy(drive)), approvals=AutoApprovals(False, by="mara.okonjo"))
    res = await a.harness.run(REQUEST)
    assert res.status == "finished"
    assert "not approved" in res.answer
    assert (drive / "Q3_Renewals.xlsx").read_bytes() == before
    assert not (drive / "Q3_Renewal_Pack.xlsx").exists()
    results = [e.data["content"] for e in a.log.of_type(L.TOOL_RESULT) if e.data["name"] in ("edit_workbook", "create_workbook")]
    assert all(r.startswith("Not written") and "mara.okonjo" in r for r in results)


async def test_permissions_are_enforced_not_suggested(drive, tmp_path):
    outside = tmp_path / "elsewhere.txt"
    outside.write_text("secret")
    policy_calls = iter([
        call(("read_file", {"path": str(drive / "HR" / "Payroll_2026_confidential.xlsx")}),
             ("read_file", {"path": str(outside)}),
             ("read_file", {"path": str(drive / "Q3_Renewals.xlsx")})),
        say("done"),
    ])
    cfg = base_config(drive, tmp_path)
    a = await assemble(cfg, model=scripted_client(lambda m, t: next(policy_calls)), approvals=AutoApprovals(True))
    await a.harness.run("read these")
    denied = a.log.of_type(L.PERMISSION_DENIED)
    assert [d.data["reason"] for d in denied] == ["excluded_by_policy", "outside_granted_folders"]
    results = a.log.of_type(L.TOOL_RESULT)
    assert results[0].data["is_error"] and "Permission denied" in results[0].data["content"]
    assert not results[2].data["is_error"]


def test_policy_exclusions_are_not_grantable(tmp_path):
    b = PermissionBroker(policy=Policy(excluded_paths=["**/HR/**"], excluded_windows=["Password manager"],
                                       disabled_grants=["input"]))
    import pytest
    from ondo_agent.permissions import PermissionDenied

    with pytest.raises(PermissionDenied):
        b.grant("files", [str(tmp_path / "HR")])
    with pytest.raises(PermissionDenied):
        b.grant("screen", ["1Password — Password manager"])
    with pytest.raises(PermissionDenied):
        b.grant("input")
    b.grant("files", [str(tmp_path)])
    assert b.check_path(tmp_path / "a.xlsx").allowed
    assert b.check_path(tmp_path / "HR" / "x.xlsx").reason == "excluded_by_policy"


def test_symlink_cannot_escape_a_granted_folder(tmp_path):
    granted = tmp_path / "granted"
    granted.mkdir()
    secret = tmp_path / "secret.txt"
    secret.write_text("x")
    (granted / "link.txt").symlink_to(secret)
    b = PermissionBroker({"files": Grant("files", True, [str(granted)])})
    assert b.check_path(granted / "link.txt").reason == "outside_granted_folders"


async def test_revoking_files_mid_run_stops_file_tools(drive, tmp_path):
    cfg = base_config(drive, tmp_path)
    state = {"a": None}

    def policy(messages, tools):
        turns = sum(1 for m in messages if m.role == "assistant")
        if turns == 0:
            return call(("list_folder", {"path": str(drive)}))
        if turns == 1:
            state["a"].broker.revoke("files", by="admin@northwind", reason="revoked from console")
            return call(("read_file", {"path": str(drive / "Q3_Renewals.xlsx")}))
        return say("stopped reading")

    a = await assemble(cfg, model=scripted_client(policy), approvals=AutoApprovals(True))
    state["a"] = a
    await a.harness.run("look")
    last = a.log.of_type(L.TOOL_RESULT)[-1]
    assert last.data["is_error"] and "not active" in last.data["content"]
    change = a.log.of_type(L.GRANT_CHANGED)[0]
    assert change.data["kind"] == "files" and change.data["by"] == "admin@northwind"


async def test_kill_switch_stops_the_run(drive, tmp_path):
    cfg = base_config(drive, tmp_path)
    holder = {}

    def policy(messages, tools):
        holder["a"].broker.stop("Escape pressed twice", by="mara.okonjo")
        return call(("list_folder", {"path": str(drive)}))

    a = await assemble(cfg, model=scripted_client(policy), approvals=AutoApprovals(True))
    holder["a"] = a
    res = await a.harness.run("anything")
    assert res.status == "stopped" and res.reason == "Escape pressed twice"
    stopped = a.log.of_type(L.RUN_STOPPED)[-1]
    assert stopped.source == "broker:mara.okonjo"
